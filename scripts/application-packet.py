"""Private profile + role draft -> local submission files. No network or model calls."""
import argparse
import hashlib
import json
import re
import sys
import time
from datetime import date, datetime
from pathlib import Path
from xml.sax.saxutils import escape


def require(condition, message):
    if not condition:
        raise ValueError(message)


def text(value):
    require(isinstance(value, str) and value.strip(), "Expected nonempty text.")
    require(not re.search(r"\[(?:TODO|TBD|insert)[^\]]*\]|\{\{.*?\}\}", value, re.I),
            "Replace draft placeholders before rendering.")
    return value.strip().translate(str.maketrans({"\u2013": "-", "\u2014": "-", "\u2011": "-"}))


def numbers(value):
    units = {word: number for number, word in enumerate(
        "zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen seventeen eighteen nineteen".split())}
    tens = {word: number * 10 for number, word in enumerate(
        "twenty thirty forty fifty sixty seventy eighty ninety".split(), 2)}
    words = "|".join([*units, *tens, "hundred", "thousand", "million", "billion", "trillion"])

    def normalize(match):
        parts = re.split(r"[\s-]+", match[0])
        if len(parts) == 1 and parts[0] in units | tens:
            return str((units | tens)[parts[0]])
        if len(parts) == 2 and parts[0] in tens and parts[1] in units and 1 <= units[parts[1]] <= 9:
            return str(tens[parts[0]] + units[parts[1]])
        raise ValueError("Use digits for unsupported number-word forms before checking evidence.")

    value = re.sub(r"\b(?:" + words + r")\b(?:[\s-]+(?:and[\s-]+)?(?:" + words + r")\b)*",
                   normalize, value.lower().replace("\u2011", "-"))
    return set(re.findall(r"(?<![\w])\$?\d[\d,]*(?:\.\d+)?(?:%|\+)?", value))


def month(value):
    require(isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}", value), "Career dates must use YYYY-MM.")
    return datetime.strptime(value, "%Y-%m")


def prepare(profile, draft):
    require(profile.get("candidate_id") and draft.get("candidate_id") == profile["candidate_id"], "Candidate identity mismatch.")
    require(profile.get("version") and draft.get("profile_version") == profile["version"], "Refresh the draft for this profile version.")
    identity = profile["identity"]
    for key in ("name", "location", "email"):
        text(identity[key])
    links = identity.get("links", [])
    for link in links:
        text(link["label"])
        require(re.match(r"^https://[^\s]+$", link["url"]), "Candidate links require HTTPS URLs.")
    sources = {s["id"]: s for s in profile["sources"]}
    require(len(sources) == len(profile["sources"]) and sources, "Source IDs must be unique.")
    for source in sources.values():
        text(source["reference"])
        text(source["revision"])
    facts = profile["facts"]
    for fact in facts.values():
        text(fact["text"])
        require(fact["source"] in sources, "A fact has no recorded source.")

    def claim(value, role=None):
        wording = text(value["text"])
        evidence = value.get("evidence", [])
        require(evidence and all(k in facts for k in evidence), "Every career claim needs known evidence IDs.")
        if role:
            require(all(facts[k].get("role") == role for k in evidence), "A bullet attributes evidence to the wrong role.")
        supported = set().union(*(numbers(facts[k]["text"]) for k in evidence))
        require(numbers(wording) <= supported, "A draft introduces an unsupported number or metric.")
        return wording

    job = draft["job"]
    for key in ("id", "company", "role"):
        text(job[key])
    require(re.match(r"^https://[^\s]+$", job["source_url"]), "Record the employer posting URL.")
    date.fromisoformat(job["checked_on"])
    canonical = {r["id"]: r for r in profile["roles"]}
    require(len(canonical) == len(profile["roles"]), "Career role IDs must be unique.")
    for role in canonical.values():
        text(role["employer"])
        text(role["title"])
        start = month(role["start"])
        require(not role.get("end") or month(role["end"]) >= start, "A career period ends before it starts.")
    resume = draft["resume"]
    require(set(resume) <= {"headline", "summary", "roles", "skills", "omitted_roles"}, "Resume fields cannot override source career facts.")
    selected, selected_ids = [], set()
    for choice in resume["roles"]:
        require(set(choice) <= {"id", "bullets"}, "Role identity and dates come from the profile, not the draft.")
        role_id = choice["id"]
        require(role_id in canonical and role_id not in selected_ids, "Unknown or duplicate selected role.")
        selected_ids.add(role_id)
        role = dict(canonical[role_id])
        role["bullets"] = [claim(b, role_id) for b in choice.get("bullets", [])]
        selected.append(role)
    require(selected, "Select at least one source-backed career role.")
    omitted = resume.get("omitted_roles", [])
    omitted_ids = {r["id"] for r in omitted}
    require(len(omitted_ids) == len(omitted) and omitted_ids == set(canonical) - selected_ids,
            "Record a reason for every omitted career role; do not silently drop history.")
    for omitted_role in omitted:
        text(omitted_role["reason"])
    selected.sort(key=lambda r: (r.get("end") or "9999-12", r["start"]), reverse=True)
    approved_skills = set(profile.get("skills", []))
    skills = [text(s) for s in resume.get("skills", [])]
    require(set(skills) <= approved_skills, "A draft introduces an unconfirmed skill.")
    requirements = {r["id"] for r in job.get("requirements", [])}
    matches = draft.get("matches", [])
    require(len({m["requirement"] for m in matches}) == len(matches), "Requirement dispositions must be unique.")
    require({m["requirement"] for m in matches} == requirements, "Map each selected job requirement or record its gap.")
    for match in matches:
        evidence = match.get("evidence", [])
        require(all(k in facts for k in evidence), "A requirement match cites unknown evidence.")
        require(evidence or match.get("gap"), "A requirement needs evidence or an explicit gap.")
    letter = draft.get("cover_letter")
    if job.get("cover_letter_required"):
        require(letter, "This employer requires a cover letter.")
    if letter:
        date.fromisoformat(letter["date"])
        letter = {**letter, "opening": text(letter["opening"]),
                  "paragraphs": [claim(p) for p in letter["paragraphs"]], "closing": text(letter["closing"])}
        require(not numbers(letter["opening"] + " " + letter["closing"]),
                "Keep measurable career claims in the evidence-backed letter paragraphs.")
    return {"identity": identity, "headline": text(resume["headline"]), "summary": claim(resume["summary"]),
            "roles": selected, "skills": skills, "education": profile.get("education", []),
            "certifications": profile.get("certifications", []), "letter": letter, "job": job,
            "unanswered": draft.get("unanswered", []), "sources": list(sources.values()), "omitted_roles": omitted}


def blocks(packet, kind):
    """One content stream drives both output formats; renderers own only presentation."""
    identity = packet["identity"]
    result = [("title", identity["name"]), ("headline", packet["headline"]), ("contact", identity)]
    if kind == "resume":
        result += [("heading", "Summary"), ("body", packet["summary"]), ("heading", "Experience")]
        for role in packet["roles"]:
            dates = month(role["start"]).strftime("%B %Y") + " - " + (month(role["end"]).strftime("%B %Y") if role.get("end") else "Present")
            group = [("role", role["employer"] + " | " + role["title"]), ("meta", dates + (" | " + role["location"] if role.get("location") else ""))]
            if role.get("context"):
                group.append(("body", role["context"]))
            group += [("bullet", b) for b in role["bullets"]]
            result.append(("group", group))
        for heading, entries in (("Skills", [" | ".join(packet["skills"])] if packet["skills"] else []),
                                 ("Education", packet["education"]), ("Certifications", packet["certifications"])):
            if entries:
                result += [("heading", heading)] + [("body", text(e)) for e in entries]
    else:
        letter = packet["letter"]
        result += [("body", date.fromisoformat(letter["date"]).strftime("%B %d, %Y")),
                   ("body", letter.get("recipient", packet["job"]["company"] + " Hiring Team")),
                   ("body", "Re: " + packet["job"]["role"]), ("body", "Dear Hiring Team,"),
                   ("body", letter["opening"])]
        result += [("body", p) for p in letter["paragraphs"]]
        result += [("body", letter["closing"]), ("body", "Best,"), ("body", identity["name"])]
    return result


def contact_parts(identity):
    parts = [(identity["location"], None)]
    if identity.get("phone"):
        parts.append((identity["phone"], None))
    parts.append((identity["email"], "mailto:" + identity["email"]))
    parts += [(link["label"], link["url"]) for link in identity.get("links", [])]
    return parts


def render_pdf(content, destination, name, kind):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, KeepTogether, Spacer
    sizes = {"title": (22, 25), "headline": (11, 14), "contact": (9, 12), "heading": (11, 14),
             "role": (10.5, 13), "meta": (9, 12), "body": (10.3 if kind == "resume" else 11, 13.5 if kind == "resume" else 15), "bullet": (10.3, 13.5)}
    styles = {key: ParagraphStyle(key, fontName="Helvetica-Bold" if key in {"title", "heading", "role"} else "Helvetica",
              fontSize=size, leading=leading, alignment=TA_LEFT, spaceAfter=4 if kind == "resume" else 10,
              spaceBefore=9 if key == "heading" else 0, keepWithNext=key in {"title", "headline", "heading", "role", "meta"},
              textColor=colors.HexColor("#26323d"), leftIndent=11 if key == "bullet" else 0,
              firstLineIndent=0, bulletIndent=1, allowWidows=0, allowOrphans=0) for key, (size, leading) in sizes.items()}

    def flow(item):
        key, value = item
        if key == "group":
            return KeepTogether([flow(child) for child in value] + [Spacer(1, 5)])
        if key == "contact":
            value = " | ".join('<link href="' + escape(url, {'"': '&quot;'}) + '">' + escape(label) + '</link>' if url else escape(label)
                               for label, url in contact_parts(value))
        else:
            value = escape(text(value))
        return Paragraph(value, styles[key], bulletText="\u2022" if key == "bullet" else None)

    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#617080"))
        if document.page > 1:
            canvas.drawString(50, 766, name + " | " + ("Resume" if kind == "resume" else "Cover letter"))
        canvas.drawRightString(562, 27, str(document.page))

    SimpleDocTemplate(str(destination), pagesize=(612, 792), leftMargin=50, rightMargin=50, topMargin=45,
                      bottomMargin=43, title=name + " " + kind, author=name).build([flow(b) for b in content], onFirstPage=footer, onLaterPages=footer)


def render_docx(content, destination, kind):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.625), Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.7)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor.from_string("26323D")
        style.font.size = Pt(10.3 if kind == "resume" else 11)
        style.paragraph_format.space_after = Pt(4 if kind == "resume" else 10)
        style.paragraph_format.line_spacing = Pt(13.5 if kind == "resume" else 15)
        style.paragraph_format.widow_control = True
    document.styles["Title"].font.size = Pt(22)
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Heading 1"].font.size = Pt(11)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(9)
    document.styles["Heading 2"].font.size = Pt(10.5)
    document.styles["List Bullet"].paragraph_format.left_indent = Inches(0.16)
    document.styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.14)
    styles = {"title": "Title", "headline": "Subtitle", "heading": "Heading 1", "role": "Heading 2", "bullet": "List Bullet"}

    def add(item, keep_next=False):
        key, value = item
        if key == "group":
            for i, child in enumerate(value):
                add(child, i < len(value) - 1)
            return
        paragraph = document.add_paragraph(style=styles.get(key, "Normal"))
        paragraph.paragraph_format.keep_with_next = keep_next or key in {"title", "headline", "heading", "role", "meta"}
        paragraph.paragraph_format.keep_together = True
        if key == "contact":
            for i, (label, url) in enumerate(contact_parts(value)):
                if i:
                    paragraph.add_run(" | ")
                if url:
                    link = OxmlElement("w:hyperlink")
                    link.set(qn("r:id"), document.part.relate_to(url, RT.HYPERLINK, is_external=True))
                    run = OxmlElement("w:r")
                    node = OxmlElement("w:t")
                    node.text = label
                    run.append(node)
                    link.append(run)
                    paragraph._p.append(link)
                else:
                    paragraph.add_run(label)
            for run in paragraph.runs:
                run.font.size = Pt(9)
        else:
            run = paragraph.add_run(text(value))
            if key == "meta":
                run.italic = True
                run.font.size = Pt(9)
    for item in content:
        add(item)
    document.core_properties.author = content[0][1]
    document.core_properties.title = content[0][1] + (" Resume" if kind == "resume" else " Cover Letter")
    document.save(destination)


def flat_text(content):
    result = []
    for key, value in content:
        if key == "group":
            result += flat_text(value)
        elif key == "contact":
            result += [label for label, _ in contact_parts(value)]
        else:
            result.append(text(value))
    return result


def check_file(path, expected, identity, page_limit):
    from pypdf import PdfReader
    from docx import Document
    normalize = lambda value: re.sub(r"\s+", "", value)
    if path.suffix == ".pdf":
        reader = PdfReader(path)
        extracted = "\n".join(p.extract_text() for p in reader.pages)
        require(0 < len(reader.pages) <= page_limit, "The document exceeds its page budget; shorten the draft instead of shrinking the type.")
        urls = {str(a.get_object().get("/A", {}).get("/URI", "")) for p in reader.pages for a in p.get("/Annots", [])}
        pages = len(reader.pages)
    else:
        document = Document(path)
        extracted = "\n".join(p.text for p in document.paragraphs)
        urls = {r.target_ref for r in document.part.rels.values() if r.reltype.endswith("/hyperlink")}
        pages = None
    require(all(normalize(value) in normalize(extracted) for value in expected), "Exported text is missing content.")
    required_urls = {url for _, url in contact_parts(identity) if url}
    require(required_urls <= urls, "Exported contact links are incomplete.")
    return {"text": "pass", "links": "pass", "pages": pages,
            "layout_review": "pending" if pages else "unverified", "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def build(profile, draft, output, formats=("pdf", "docx")):
    started = time.perf_counter()
    packet = prepare(profile, draft)
    output = Path(output)
    require(not output.exists() or not any(output.iterdir()), "Use a new or empty output folder; existing packet files are preserved.")
    require(set(formats) <= {"pdf", "docx"} and formats, "Formats must be pdf and/or docx.")
    output.mkdir(parents=True, exist_ok=True)
    report = {"stage": "review_required", "profile_version": profile["version"], "job_id": packet["job"]["id"],
              "sources": packet["sources"], "omitted_roles": packet["omitted_roles"], "unanswered": packet["unanswered"],
              "checks": {"candidate_binding": "pass", "evidence_references": "pass", "numeric_claims": "pass",
                         "career_identity": "source_preserved", "chronology": "source_dates_sorted", "requirements": "dispositions_recorded"},
              "review_required": ["Claims preserve source meaning", "Role fit and candidate voice", "Every final PDF page",
                                  "DOCX layout if DOCX will be submitted"], "files": {}}
    for kind in ("resume", "cover-letter"):
        if kind == "cover-letter" and not packet["letter"]:
            continue
        content = blocks(packet, kind)
        for fmt in formats:
            path = output / (kind + "." + fmt)
            if fmt == "pdf":
                render_pdf(content, path, packet["identity"]["name"], kind)
            else:
                render_docx(content, path, kind)
            report["files"][path.name] = check_file(path, flat_text(content), packet["identity"],
                                                  packet["job"].get("resume_page_limit", 2) if kind == "resume" else 1)
            if fmt == "pdf":
                import pypdfium2 as pdfium
                pdf = pdfium.PdfDocument(path)
                preview_dir = output / "preview"
                preview_dir.mkdir(exist_ok=True)
                for i in range(len(pdf)):
                    page = pdf[i]
                    bitmap = page.render(scale=1.6)
                    bitmap.to_pil().save(preview_dir / (kind + "-" + str(i + 1) + ".png"))
                    bitmap.close()
                    page.close()
                pdf.close()
    report["mechanical_seconds"] = round(time.perf_counter() - started, 3)
    (output / "build-report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--profile", type=Path, required=True)
    parser.add_argument("--draft", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--formats", default="pdf,docx")
    args = parser.parse_args()
    try:
        profile, draft = [json.loads(p.read_text(encoding="utf-8-sig")) for p in (args.profile, args.draft)]
        report = build(profile, draft, args.output, tuple(args.formats.split(",")))
        print(json.dumps({"stage": report["stage"], "mechanical_seconds": report["mechanical_seconds"],
                          "files": {name: {"pages": data["pages"], "text": data["text"], "links": data["links"]}
                                    for name, data in report["files"].items()}, "unanswered_count": len(report["unanswered"])}))
    except (ValueError, KeyError, TypeError, OSError, ImportError) as error:
        print(json.dumps({"stage": "needs_attention", "reason": str(error)}), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
